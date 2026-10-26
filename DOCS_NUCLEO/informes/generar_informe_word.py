# -*- coding: utf-8 -*-
"""
Script generador del Informe Técnico de Trabajo en formato Word (.docx)
NÚCLEO PRINCIPAL - POI 2026
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generar_informe_docx(ruta_salida):
    doc = docx.Document()

    # Configuración de Márgenes
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Paleta de colores institucionales
    COLOR_AZUL = RGBColor(16, 38, 56)      # #102638
    COLOR_CIAN = RGBColor(0, 142, 214)     # #008ED6

    # Título Principal
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run('INFORME TÉCNICO DE TRABAJO DIARIO')
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_AZUL

    # Subtítulo
    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run('PROYECTO: VISOR HIDROMETEOROLÓGICO CUENCA DEL RÍO TUMBES (POI 2026)')
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(10.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_CIAN

    # Tabla de Metadatos
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ('Organización:', 'Dirección de Hidrología Aplicada / SENAMHI - POI 2026'),
        ('Fecha de Ejecución:', '26 de Octubre de 2026'),
        ('Responsables:', 'Ing. Max Gutierrez & NÚCLEO PRINCIPAL'),
        ('Estado Operativo:', 'Hito 1 Concluido - 100% Operativo, Automatizado y en la Nube')
    ]
    for idx, (k, v) in enumerate(meta_data):
        row = meta_table.rows[idx]
        c0 = row.cells[0].paragraphs[0].add_run(k)
        c0.font.name = 'Arial'
        c0.font.bold = True
        c0.font.size = Pt(9.5)
        c1 = row.cells[1].paragraphs[0].add_run(v)
        c1.font.name = 'Arial'
        c1.font.size = Pt(9.5)
        if idx == 3:
            c1.font.bold = True
            c1.font.color.rgb = RGBColor(0, 128, 64)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def agregar_seccion(texto):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(texto)
        r.font.name = 'Arial'
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_AZUL
        return p

    # 1. OBJETIVO GENERAL
    agregar_seccion('1. OBJETIVO GENERAL')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(
        'Auditar, refactorizar, optimizar y automatizar de extremo a extremo el sistema '
        'hidro-meteorológico del Visor TUMBES, subsanando fallas críticas de rutas fijas y dependencias, '
        'protegiendo credenciales institucionales, garantizando la portabilidad del código y desplegando '
        'el pipeline continuo de ingesta satelital en GitHub Actions sincronizado con la aplicación interactiva '
        'en Streamlit Cloud.'
    )
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)

    # 2. DIAGNÓSTICO TÉCNICO INICIAL
    agregar_seccion('2. DIAGNÓSTICO TÉCNICO INICIAL')
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run('Al comenzar la intervención, el repositorio presentaba deficiencias estructurales que impedían su ejecución fuera de la estación de origen:')
    r.font.name = 'Arial'
    r.font.size = Pt(9.5)

    puntos_diag = [
        ('Rutas Absolutas Hardcodeadas:', ' Todos los scripts tenían quemada la unidad E:\\tumbes_262\\..., colapsando al ejecutarse en cualquier otra computadora o servidor.'),
        ('Exposición de Credenciales en Texto Plano:', ' Contraseñas de acceso a la NASA PPS y tokens de SENAMHI se encontraban dentro de los scripts.'),
        ('Gráfica de Caudal Desconectada:', ' La Sección 2 del visor en Streamlit se encontraba vacía como un marcador provisional (placeholder).'),
        ('Riesgo de Saturación en la Nube:', ' Acumulación continua de rasters satelitales GeoTIFF sin política de depuración temporal.'),
        ('Incompatibilidad de Dependencias C/GDAL:', ' Ausencia de drivers y manifiesto para runners Linux.')
    ]
    for tit, desc in puntos_diag:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_after = Pt(2)
        r1 = p_b.add_run(tit)
        r1.font.name = 'Arial'
        r1.font.bold = True
        r1.font.size = Pt(9)
        r2 = p_b.add_run(desc)
        r2.font.name = 'Arial'
        r2.font.size = Pt(9)

    # 3. ACTIVIDADES EJECUTADAS Y LOGROS
    agregar_seccion('3. ACTIVIDADES EJECUTADAS Y LOGROS ALCANZADOS')

    acts = [
        ('3.1 Portabilidad Dinámica y Desacople Seguro', 'Creación de backend/codes/config.py implementando pathlib.Path para resolución relativa. Desacople de claves hacia archivo .env local y secretos cifrados en GitHub Secrets.'),
        ('3.2 Validación Operativa Hidrológica en Tiempo Real', 'Descarga exitosa de las últimas 72 horas de caudal en la estación El Tigre vía API SENAMHI (55.1 m³/s - Condición Normal). Extracción zonal de 172 rasters GeoTIFF NASA GPM IMERG Early sobre la cuenca del Río Tumbes y muestreo en las 13 estaciones pluviométricas de la red.'),
        ('3.3 Conexión de Caudal y Umbrales Oficiales en Streamlit', 'Integración de serie horaria real en frontend/app.py incorporando las líneas de alerta oficial de SENAMHI: Amarillo (600 m³/s), Naranja (1,000 m³/s) y Rojo (1,500 m³/s), con cálculo de picos y visualización en tiempo real.'),
        ('3.4 Política de Retención y Control de Peso (Anti-Saturación)', 'Desarrollo de backend/codes/limpieza_retencion.py implementando una ventana deslizante de 72 horas. Purga automática de datos históricos que mantiene el repositorio en un peso constante de ~10 a 15 MB.'),
        ('3.5 Automatización en la Nube (GitHub Actions - Estilo AMARU)', 'Diseño del workflow .github/workflows/actualizacion_automatica.yml programado cada 3 horas (cron 0 */3 * * *) para compilar en Ubuntu, descargar datos, correr el orquestador nativo y hacer push automático.'),
        ('3.6 Estandarización para Despliegue en Streamlit Cloud', 'Creación de packages.txt (libgdal-dev, gdal-bin) y optimización de requirements.txt eliminando más de 180 paquetes innecesarios para garantizar compilación limpia en servidores Linux.'),
        ('3.7 Infraestructura Documental NÚCLEO', 'Creación de DOCS_NUCLEO/ con 6 reportes analíticos, 11 bitácoras de avance y plantillas exportables para LLMs (OpenAI Codex / Claude).')
    ]

    for tit, desc in acts:
        p_a = doc.add_paragraph()
        p_a.paragraph_format.space_before = Pt(3)
        p_a.paragraph_format.space_after = Pt(2)
        r_a1 = p_a.add_run(tit + ': ')
        r_a1.font.name = 'Arial'
        r_a1.font.bold = True
        r_a1.font.size = Pt(9)
        r_a1.font.color.rgb = COLOR_CIAN
        r_a2 = p_a.add_run(desc)
        r_a2.font.name = 'Arial'
        r_a2.font.size = Pt(9)

    # 4. MATRIZ DE ESTADO OPERATIVO
    agregar_seccion('4. MATRIZ DE ESTADO OPERATIVO DE COMPONENTES')

    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = tbl.rows[0].cells
    headers = ['Componente / Script', 'Función Operativa', 'Estado']
    for idx, text in enumerate(headers):
        cell = hdr_cells[idx]
        shd = parse_xml(r'<w:shd {} w:fill="102638"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shd)
        r = cell.paragraphs[0].add_run(text)
        r.font.name = 'Arial'
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    filas = [
        ('config.py', 'Detección dinámica de rutas y variables de entorno', '✅ 100% Operativo'),
        ('descarga_q_tigre.py', 'Descarga de caudal horario en estación El Tigre', '✅ 100% Operativo'),
        ('extraccion_areal.py', 'Cálculo zonal de precipitación satelital en cuenca', '✅ 100% Operativo'),
        ('extraer_lluvia_estaciones.py', 'Muestreo puntual en 13 estaciones pluviométricas', '✅ 100% Operativo'),
        ('limpieza_retencion.py', 'Ventana rodante de 72h (Control de peso)', '✅ 100% Operativo'),
        ('orquestador_tumbes.py', 'Orquestador maestro secuencial del autor', '✅ 100% Operativo'),
        ('actualizacion_automatica.yml', 'GitHub Actions cada 3 horas en la nube', '✅ Activo en GitHub'),
        ('frontend/app.py', 'Dashboard Streamlit con mapa y caudal real', '✅ Conectado'),
        ('packages.txt / requirements.txt', 'Drivers GDAL y librerías de producción', '✅ Resuelto')
    ]

    for c, f, s in filas:
        row_cells = tbl.add_row().cells
        r0 = row_cells[0].paragraphs[0].add_run(c)
        r0.font.name = 'Arial'
        r0.font.bold = True
        r0.font.size = Pt(8)
        r1 = row_cells[1].paragraphs[0].add_run(f)
        r1.font.name = 'Arial'
        r1.font.size = Pt(8)
        r2 = row_cells[2].paragraphs[0].add_run(s)
        r2.font.name = 'Arial'
        r2.font.size = Pt(8)
        r2.font.bold = True
        r2.font.color.rgb = RGBColor(0, 128, 64)

    # 5. CONCLUSIONES Y PRÓXIMOS PASOS
    agregar_seccion('5. CONCLUSIONES Y PRÓXIMOS PASOS')
    p_c = doc.add_paragraph()
    p_c.paragraph_format.space_after = Pt(3)
    rc = p_c.add_run(
        'El Visor TUMBES ha completado exitosamente su Hito 1 de modernización y portabilidad, '
        'logrando independencia de rutas fijas, protección estricta de credenciales y ejecución desatendida '
        'en la nube con bajo consumo de almacenamiento.'
    )
    rc.font.name = 'Arial'
    rc.font.size = Pt(9)

    pasos = [
        'Interacción en Mapa: Visualización de lluvia acumulada 24h/72h al interactuar con las 13 estaciones del visor.',
        'Modelación Predictiva: Desarrollo del módulo de pronóstico hidrológico en backend/model/eltigre.',
        'Sistema de Alertas Tempranas: Emisión de notificaciones push (Telegram / Webhooks) ante alertas hidrológicas.'
    ]
    for paso in pasos:
        p_p = doc.add_paragraph(style='List Bullet')
        p_p.paragraph_format.space_after = Pt(2)
        rp = p_p.add_run(paso)
        rp.font.name = 'Arial'
        rp.font.size = Pt(9)

    doc.save(ruta_salida)
    print("Documento guardado con éxito en:", ruta_salida)

if __name__ == "__main__":
    generar_informe_docx("DOCS_NUCLEO/informes/INFORME_TECNICO_TRABAJO_2026-10-26.docx")
